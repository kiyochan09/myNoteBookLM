"""Stable OCR note links and native Word footnotes."""
import json,re
from pathlib import Path
from docx.oxml import OxmlElement,parse_xml
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

TOKEN=re.compile(r'⟦注:([a-zA-Z0-9_-]+)⟧')

def note_id(note,index):
    return str(note.get('note_id') or note.get('region_id') or note.get('id') or f'note-{index}') if isinstance(note,dict) else f'note-{index}'

def note_text(note):
    return str(note.get('text','') if isinstance(note,dict) else note)

def build_catalog(directory):
    pages={}
    for path in sorted(Path(directory).glob('page_*/page_data.json')):
        try:
            data=json.loads(path.read_text(encoding='utf-8'))
            page=int(data.get('page_number') or data.get('current_page') or path.parent.name.split('_')[-1])
            pages[str(page)]={'body_text':data.get('body_text',''),'footnote_links':data.get('footnote_links',[]),
                'footnotes':[{'note_id':note_id(n,i),'text':note_text(n)} for i,n in enumerate(data.get('footnotes',[]))]}
        except (ValueError,OSError,TypeError):
            continue
    return pages

class WordFootnotes:
    def __init__(self,document,page_data,directory,memory_pages=None):
        self.document=document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt
        if 'Footnote Text' not in document.styles:
            style=document.styles.add_style('Footnote Text',WD_STYLE_TYPE.PARAGRAPH)
            style.font.size=Pt(9)
        if 'Footnote Reference' not in document.styles:
            style=document.styles.add_style('Footnote Reference',WD_STYLE_TYPE.CHARACTER)
            style.font.superscript=True
        self.directory=Path(directory)
        self.pages={int(n):d for n,d,_ in page_data}
        for data in memory_pages or []:
            n=data.get('page_number') or data.get('page') or data.get('current_page')
            if n:self.pages[int(n)]=data
        self.links={l['id']:l for _,data,_ in page_data for l in data.get('footnote_links',[]) if l.get('id')}
        self.bound=set()
        self.root=OxmlElement('w:footnotes')
        for identity,kind in [(-1,'separator'),(0,'continuationSeparator')]:
            note=OxmlElement('w:footnote');note.set(qn('w:id'),str(identity));note.set(qn('w:type'),kind)
            p=OxmlElement('w:p');r=OxmlElement('w:r');r.append(OxmlElement('w:'+kind));p.append(r);note.append(p);self.root.append(note)
        self.part=None
        self.next_id=1
        # Suppress the former end-of-document duplicate only for linked targets
        # that actually occur in the exported body.
        for _,data,_ in page_data:
            for match in TOKEN.finditer(data.get('body_text','')):
                link=self.links.get(match.group(1))
                if link and link.get('note_page') and link.get('note_id'):
                    self.bound.add((int(link['note_page']),str(link['note_id'])))

    def resolve(self,identity):
        link=self.links.get(identity)
        if not link or not link.get('note_page') or not link.get('note_id'):
            raise ValueError('注釈番号に対応する注釈文を選択してからDOCXを出力してください。')
        page=int(link['note_page']);data=self.pages.get(page)
        if data is None:
            path=self.directory/f'page_{page:04d}'/'page_data.json'
            if path.exists():data=json.loads(path.read_text(encoding='utf-8'));self.pages[page]=data
        if data:
            for index,note in enumerate(data.get('footnotes',[])):
                if note_id(note,index)==str(link['note_id']):return note_text(note)
        raise ValueError('関連付けた注釈文が見つかりません。本文の注釈番号を設定し直してください。')

    def write(self,paragraph,text):
        position=0
        for match in TOKEN.finditer(text):
            paragraph.add_run(text[position:match.start()])
            content=self.resolve(match.group(1));identity=self.next_id;self.next_id+=1
            if self.part is None:
                self.part=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',b'',self.document.part.package)
                self.document.part.relate_to(self.part,RT.FOOTNOTES)
            reference=OxmlElement('w:footnoteReference');reference.set(qn('w:id'),str(identity))
            run=paragraph.add_run();run.style='Footnote Reference'
            from docx.shared import RGBColor
            run.font.color.rgb=RGBColor(0x16,0xa3,0x4a)
            run.font.superscript=True;run._r.append(reference)
            note=OxmlElement('w:footnote');note.set(qn('w:id'),str(identity))
            p=OxmlElement('w:p');properties=OxmlElement('w:pPr');style=OxmlElement('w:pStyle');style.set(qn('w:val'),'FootnoteText');properties.append(style);p.append(properties)
            r=OxmlElement('w:r');properties=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'FootnoteReference');properties.append(style);r.append(properties);r.append(OxmlElement('w:footnoteRef'));p.append(r)
            for i,line in enumerate(content.split('\n')):
                r=OxmlElement('w:r')
                if i:r.append(OxmlElement('w:br'))
                t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=(' ' if i==0 else '')+line;r.append(t);p.append(r)
            note.append(p);self.root.append(note)
            self.part._blob=__import__('lxml.etree',fromlist=['tostring']).tostring(self.root,xml_declaration=True,encoding='UTF-8',standalone=True)
            position=match.end()
        paragraph.add_run(text[position:])
