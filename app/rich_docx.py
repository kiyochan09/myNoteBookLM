"""Serialize editor HTML, keeping source formatting where text still matches."""
import copy
import io
import re
import difflib
from collections import defaultdict
from pathlib import Path
from app.paths import AppPaths

import docx
from lxml import etree, html
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


def css(node):
    return {k.strip().lower(): v.strip() for item in node.get('style', '').split(';') if ':' in item for k, v in [item.split(':', 1)]}


def points(value, em=11):
    m = re.fullmatch(r'\s*(-?[\d.]+)(pt|px|em|rem|in|cm|mm)?\s*', str(value))
    if not m:
        return None
    return float(m[1]) * {'pt': 1, 'px': .75, 'em': em, 'rem': 12, 'in': 72, 'cm': 72 / 2.54, 'mm': 72 / 25.4, None: 1}[m[2]]


def color(value):
    value = value.strip().lower()
    named = {'black': '000000', 'white': 'ffffff', 'red': 'ff0000', 'blue': '0000ff', 'green': '008000', 'yellow': 'ffff00'}
    value = named.get(value, value.lstrip('#'))
    if re.fullmatch('[0-9a-f]{3}', value):
        value = ''.join(c * 2 for c in value)
    if re.fullmatch('[0-9a-f]{6}', value):
        return value.upper()
    m = re.fullmatch(r'rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,[^)]+)?\)', value)
    return ''.join(f'{min(int(c),255):02X}' for c in m.groups()) if m else None


def clean(node):
    node = copy.deepcopy(node)
    for x in node.xpath('.//*[contains(concat(" ",normalize-space(@class)," ")," fn-anchor ")]'):
        x.drop_tree()
    for x in node.xpath('.//br'):
        x.tail = '\n' + (x.tail or '')
    return ''.join(node.itertext())


def word_text(node):
    return ''.join((x.text or '') if x.tag == qn('w:t') else '\u2003' if x.tag == qn('w:tab') else '\n' if x.tag in (qn('w:br'), qn('w:cr')) else '' for x in node.iter())


class RichDocx:
    def __init__(self, bundle, base_dir):
        self.bundle = bundle
        self.base_dir = Path(base_dir)
        self.source = defaultdict(list)
        self.source_order = []
        self.source_by_block = {}
        self.doc = docx.Document()
        name = bundle['document'].get('source_filename')
        source = self.base_dir / 'data' / Path(name or '').name
        if name and source.suffix.lower() == '.docx' and source.is_file():
            self.doc = docx.Document(source)
            for p in self.doc.element.body.iter(qn('w:p')):
                text = word_text(p)
                self.source[text].append(copy.deepcopy(p))
                self.source_order.append((text, copy.deepcopy(p)))
            # Keep source styles, defaults, numbering and related parts; replace body.
            for element in list(self.doc.element.body):
                if element.tag != qn('w:sectPr'):
                    self.doc.element.body.remove(element)
        targets = []
        for block in bundle['blocks']:
            try:
                nodes = html.fragment_fromstring(block.get('html_content') or '<p></p>', create_parent='div')
                if len(nodes) == 1 and re.fullmatch(r'p|h[1-6]|blockquote', nodes[0].tag):
                    targets.append((block['id'], clean(nodes[0])))
            except (etree.ParserError, ValueError):
                continue
        matcher = difflib.SequenceMatcher(None, [t for t, _ in self.source_order], [t for _, t in targets], autojunk=False)
        for source_i, target_i, count in matcher.get_matching_blocks():
            for j in range(count):
                self.source_by_block[targets[target_i+j][0]] = self.source_order[source_i+j][1]
        self.notes = bundle.get('annotations', [])
        self.note_keys = {}
        self.used = set()
        self.note_roots = {}
        self.next_note = defaultdict(lambda: 1)
        self.bookmark_id = 0
        self.bookmarks = {}
        self.links = []
        self.marker_notes = defaultdict(list)
        for note in self.notes:
            self.marker_notes[re.sub(r'[【】\[\]\s]', '', note.get('anchor_text', ''))].append(note)
            kind = note.get('kind')
            identity = str(note.get('id', ''))
            prefix = f'ann-{bundle["document"]["id"]}-{kind}-'
            if identity.startswith(prefix):
                self.note_keys[f'{kind}:{identity[len(prefix):]}'] = note
        self.sections = {s['id']: s for s in bundle.get('sections', [])}

    def paragraph_format(self, p, node):
        s = css(node)
        fmt = p.paragraph_format
        base_size = next((r.font.size.pt for r in p.runs if r.font.size is not None), None)
        if base_size is None:
            style = p.style
            while style is not None and base_size is None:
                if style.font.size is not None: base_size = style.font.size.pt
                style = style.base_style
        size = points(s.get('font-size', ''), base_size or 11) or base_size or 11
        align = s.get('text-align', node.get('align'))
        if align in ('left', 'center', 'right', 'justify'):
            fmt.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        for key, attr in [('margin-top', 'space_before'), ('margin-bottom', 'space_after'), ('text-indent', 'first_line_indent'), ('margin-right', 'right_indent')]:
            n = points(s.get(key, ''), size)
            if n is not None:
                setattr(fmt, attr, Pt(n))
        left = [points(s.get(k, ''), size) for k in ('margin-left', 'padding-left', 'padding-inline-start')]
        if any(v is not None for v in left):
            fmt.left_indent = Pt(sum(v or 0 for v in left))
        lh = s.get('line-height')
        if lh and lh != 'normal':
            if re.fullmatch(r'[\d.]+', lh):
                fmt.line_spacing = float(lh)
            elif lh.endswith('%'):
                fmt.line_spacing = float(lh[:-1]) / 100
            else:
                n = points(lh, size)
                if n is not None:
                    fmt.line_spacing = Pt(n)
        if s.get('page-break-before') == 'always':
            fmt.page_break_before = True
        if s.get('page-break-inside') == 'avoid':
            fmt.keep_together = True

    def run_format(self, run, styles):
        f = run.font
        if 'font-weight' in styles:
            f.bold = styles['font-weight'] in ('bold', 'bolder') or (styles['font-weight'].isdigit() and int(styles['font-weight']) >= 600)
        if 'font-style' in styles:
            f.italic = styles['font-style'] == 'italic'
        decoration = styles.get('text-decoration', styles.get('text-decoration-line'))
        if decoration is not None:
            f.underline = 'underline' in decoration
            f.strike = 'line-through' in decoration
        if 'font-size' in styles:
            n = points(styles['font-size'])
            if n is not None:
                f.size = Pt(n)
        if 'font-family' in styles:
            name = styles['font-family'].split(',')[0].strip(' "\'')
            f.name = name
            run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
        c = color(styles.get('color', ''))
        if c:
            f.color.rgb = RGBColor.from_string(c)
        c = color(styles.get('background-color', styles.get('background', '')))
        if c:
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), c)
            run._r.get_or_add_rPr().append(shade)
        if styles.get('vertical-align') in ('super', 'sub', 'baseline'):
            f.superscript = styles['vertical-align'] == 'super'
            f.subscript = styles['vertical-align'] == 'sub'

    def add_text(self, p, text, styles):
        if not text:
            return
        # Source run properties survive across HTML splits, including fonts and language.
        pos = self.offset
        end = pos + len(text)
        cuts = {pos, end}
        for a, b, _ in self.source_runs:
            if pos < a < end: cuts.add(a)
            if pos < b < end: cuts.add(b)
        ordered = sorted(cuts)
        for a, b in zip(ordered, ordered[1:]):
            r = p.add_run(text[a-pos:b-pos])
            for start, stop, props in self.source_runs:
                if start <= a < stop and props is not None:
                    r._r.insert(0, copy.deepcopy(props)); break
            self.run_format(r, styles)
        self.offset = end

    def bookmark(self, p, name):
        if name in self.bookmarks:
            return
        identity = self.bookmark_id; self.bookmark_id += 1
        safe = name
        self.bookmarks[name] = safe
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(identity)); start.set(qn('w:name'), safe)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(identity))
        p._p.append(start); p._p.append(end)

    def note(self, p, note):
        kind = note.get('kind')
        if kind not in ('footnote', 'endnote', 'comment'):
            return False
        self.used.add(note['id'])
        identity = self.next_note[kind]; self.next_note[kind] += 1
        root = self.note_roots.setdefault(kind, OxmlElement('w:' + {'footnote':'footnotes', 'endnote':'endnotes', 'comment':'comments'}[kind]))
        element = OxmlElement('w:' + kind); element.set(qn('w:id'), str(identity))
        if kind == 'comment': element.set(qn('w:author'), '注釈')
        for line_index, line in enumerate(str(note.get('target_value', '')).split('\n')):
            para = OxmlElement('w:p'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
            if kind != 'comment' and line_index == 0:
                number = OxmlElement('w:r'); number.append(OxmlElement('w:' + kind + 'Ref')); para.append(number)
            t.set(qn('xml:space'), 'preserve'); t.text = line; r.append(t); para.append(r); element.append(para)
        root.append(element)
        ref = OxmlElement('w:' + kind + 'Reference'); ref.set(qn('w:id'), str(identity))
        if kind == 'comment':
            for tag in ('commentRangeStart', 'commentRangeEnd'):
                mark = OxmlElement('w:' + tag); mark.set(qn('w:id'), str(identity)); p._p.append(mark)
        run = p.add_run(); run.font.superscript = True; run._r.append(ref)
        return True

    def inline(self, p, node, inherited=None):
        tag = node.tag.lower() if isinstance(node.tag, str) else ''
        classes = node.get('class', '').split()
        if 'fn-anchor' in classes:
            note = self.note_keys.get(node.get('data-note-key'))
            if note is None:
                candidates = self.marker_notes.get(re.sub(r'[【】\[\]\s]', '', ''.join(node.itertext())), [])
                if len(candidates) == 1: note = candidates[0]
            if note and self.note(p, note): return
        if node.get('id') and any(c in classes for c in ('doc-bookmark','bm-anchor')):
            self.bookmark(p, node.get('id'))
        styles = dict(inherited or {})
        semantics = {'b': {'font-weight':'bold'}, 'strong':{'font-weight':'bold'}, 'i':{'font-style':'italic'}, 'em':{'font-style':'italic'}, 'u':{'text-decoration':'underline'}, 's':{'text-decoration':'line-through'}, 'del':{'text-decoration':'line-through'}, 'sup':{'vertical-align':'super'}, 'sub':{'vertical-align':'sub'}, 'code':{'font-family':'Consolas'}}
        styles.update(semantics.get(tag, {}))
        local = css(node)
        if 'font-size' in local:
            parent_size = points(styles.get('font-size', '11pt')) or 11
            value = local['font-size']
            resolved = parent_size * float(value[:-1]) / 100 if re.fullmatch(r'[\d.]+%', value) else points(value, parent_size)
            if resolved is not None: local['font-size'] = str(resolved) + 'pt'
        styles.update(local)
        if tag == 'br':
            p.add_run().add_break(); self.offset += 1
            return
        start_index = len(p._p)
        if node.text: self.add_text(p, node.text, styles)
        for child_index, child in enumerate(node):
            if child_index and (child.tag in ('p','div') or css(child).get('display') == 'block'):
                p.add_run().add_break()
            self.inline(p, child, styles)
            if child.tail: self.add_text(p, child.tail, styles)
        if tag == 'a' and node.get('href'):
            href = node.get('href'); link = OxmlElement('w:hyperlink')
            if href.startswith('#'):
                link.set(qn('w:anchor'), href[1:]); self.links.append((link, href[1:]))
            elif re.match(r'^(https?://|mailto:)', href, re.I):
                rid = p.part.relate_to(href, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                link.set(qn('r:id'), rid)
            else: return
            for element in list(p._p)[start_index:]: link.append(element)
            p._p.append(link)

    def paragraph(self, node, block=None, parent=None):
        match = re.fullmatch('h([1-6])', node.tag.lower())
        level = int(match[1]) if match else None
        if level is None and block and block.get('block_type') == 'heading':
            level = max(1, min(6, int(self.sections.get(block.get('section_id'), {}).get('level', 1))))
        p = (parent or self.doc).add_paragraph()
        text = clean(node)
        originals = self.source.get(text, [])
        original = self.source_by_block.get(block.get('id')) if block else None
        if original is None and len(originals) == 1: original = originals[0]
        self.source_runs = []; self.offset = 0
        if original is not None:
            props = original.find(qn('w:pPr'))
            if props is not None: p._p.insert(0, copy.deepcopy(props))
            pos = 0
            for r in original.iter(qn('w:r')):
                rt = word_text(r)
                self.source_runs.append((pos, pos+len(rt), r.find(qn('w:rPr')))); pos += len(rt)
        if level:
            if original is None or p._p.get_or_add_pPr().find(qn('w:pStyle')) is None:
                if f'Heading {level}' not in self.doc.styles:
                    self.doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)
                p.style = f'Heading {level}'
            outline = p._p.get_or_add_pPr().find(qn('w:outlineLvl'))
            if outline is None: outline = OxmlElement('w:outlineLvl'); p._p.get_or_add_pPr().append(outline)
            outline.set(qn('w:val'), str(level-1))
        self.paragraph_format(p, node)
        self.inline(p, node)
        return p

    def finish_notes(self):
        for link, name in self.links:
            if name in self.bookmarks: link.set(qn('w:anchor'), self.bookmarks[name])
        # A missing anchor never drops note text or guesses a body position.
        remaining = [n for n in self.notes if n['id'] not in self.used]
        if remaining:
            self.doc.add_paragraph('参照位置未確定の注釈')
            for n in remaining:
                self.doc.add_paragraph(str(n.get('anchor_text') or n.get('kind', '注釈')) + '\n' + str(n.get('target_value', '')))
        for kind, root in self.note_roots.items():
            if kind != 'comment':
                for identity, typ in [(-1, 'separator'), (0, 'continuationSeparator')]:
                    note = OxmlElement('w:' + kind); note.set(qn('w:id'), str(identity)); note.set(qn('w:type'), typ)
                    p = OxmlElement('w:p'); r = OxmlElement('w:r'); r.append(OxmlElement('w:' + typ)); p.append(r); note.append(p); root.insert(0, note)
            plural = {'footnote':'footnotes', 'endnote':'endnotes', 'comment':'comments'}[kind]
            reltype = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + plural
            # Replace source definitions after rebuilding references from DB.
            for rid, rel in list(self.doc.part.rels.items()):
                if rel.reltype == reltype: del self.doc.part.rels[rid]
            part = Part(PackURI('/word/' + plural + '.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.' + plural + '+xml', etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True), self.doc.part.package)
            self.doc.part.relate_to(part, reltype)

    def export(self):
        figures = {f['id']: f for f in self.bundle.get('figures', [])}
        for block in self.bundle['blocks']:
            fragment = block.get('html_content') or '<p></p>'
            try: wrapper = html.fragment_fromstring(fragment, create_parent='div')
            except (etree.ParserError, ValueError):
                wrapper = html.Element('div'); node = html.SubElement(wrapper, 'p'); node.text = block.get('text_content', '')
            figure = wrapper.xpath('.//*[@data-fig-id]')
            if figure:
                fig = figures.get(figure[0].get('data-fig-id'), {})
                try: path = AppPaths(self.base_dir).media_file(fig.get('file_path', ''))
                except (ValueError, OSError): path = None
                if path is not None and path.is_file(): self.doc.add_picture(str(path))
                self.doc.add_paragraph(fig.get('caption', ''))
                continue
            tables = wrapper.xpath('.//table')
            if tables:
                for table in tables:
                    rows = table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
                    count = max((sum(int(c.get('colspan', '1')) for c in r.xpath('./td|./th')) for r in rows), default=0)
                    if not count: continue
                    target = self.doc.add_table(rows=len(rows), cols=count)
                    if 'Table Grid' in self.doc.styles: target.style = 'Table Grid'
                    occupied = set()
                    for ri, row in enumerate(rows):
                        ci = 0
                        for cell in row.xpath('./td|./th'):
                            while (ri, ci) in occupied: ci += 1
                            cs = int(cell.get('colspan', '1')); rs = int(cell.get('rowspan', '1'))
                            if ci+cs > count or ri+rs > len(rows): raise ValueError('表のセル結合範囲が不正です。')
                            dest = target.cell(ri, ci)
                            if cs > 1 or rs > 1: dest = dest.merge(target.cell(ri+rs-1, ci+cs-1))
                            for rr in range(ri, ri+rs):
                                for cc in range(ci, ci+cs): occupied.add((rr, cc))
                            paragraphs = cell.xpath('./p|./h1|./h2|./h3|./h4|./h5|./h6') or [cell]
                            for node in paragraphs: self.paragraph(node, parent=dest)
                            dest._tc.remove(dest.paragraphs[0]._p)
                            ci += cs
                continue
            nodes = list(wrapper)
            if not nodes:
                n = html.Element('p'); n.text = block.get('text_content', ''); nodes = [n]
            for node in nodes:
                if 'doc-textbox' in node.get('class', '').split() and node.tag == 'div':
                    for paragraph in node:
                        self.paragraph(paragraph)
                else:
                    self.paragraph(node, block)
        self.finish_notes()
        output = io.BytesIO(); self.doc.save(output); output.seek(0)
        return output


def export_rich_docx(bundle, base_dir):
    return RichDocx(bundle, base_dir).export()
